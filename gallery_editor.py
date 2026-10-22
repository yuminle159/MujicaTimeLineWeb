"""本地画廊标签编辑器：人工看图填写，只保存 tags 单元格。"""
import hashlib
import os
import shutil
import tempfile
import threading
import queue
from datetime import datetime
from pathlib import Path
import tkinter as tk
from tkinter import ttk, messagebox

import openpyxl
from PIL import Image, ImageOps, ImageTk
from generate_all import gallery_image_key, is_gallery_image, parse_gallery_tags

DIMENSIONS = ('类型', 'LIVE', '人物', '关联人物', '作品', '节目', '来源')
DEFAULT_TAGS = {'类型': ['Live返图', '访谈', '合照', '日常', '主视觉', '场刊', '唱片封面'],
                '人物': ['佐佐木李子', '高尾奏音', '渡濑结月', '冈田梦以', '米泽茜']}


def format_tags(groups):
    return ', '.join(f'{category}:{value}' if value else category
                     for category, values in groups.items() for value in (values or ['']))


def validate_tags(raw):
    """保留已有格式兼容性，同时阻止常见的中文分隔符误输入。"""
    if '，' in raw or '：' in raw:
        raise ValueError('标签请使用英文逗号和冒号，例如：类型:访谈, 人物:高尾奏音')
    for token in raw.split(','):
        if token.strip() and ':' in token:
            category, value = token.split(':', 1)
            if not category.strip() or not value.strip():
                raise ValueError('分类和标签都不能为空。')
    return format_tags(parse_gallery_tags(raw.strip()))


class GalleryStore:
    def __init__(self, root):
        self.root = Path(root).resolve()
        self.path = self.root / '_data' / 'data.xlsx'
        self.entries = []
        self.load()

    @staticmethod
    def columns(sheet):
        headers = {str(cell.value).strip().lower(): cell.column for cell in sheet[1] if cell.value}
        if not {'filename', 'tags'} <= headers.keys():
            raise ValueError('gallery_images 表需要 filename 和 tags 列。')
        return headers

    def load(self):
        workbook = openpyxl.load_workbook(self.path, read_only=True)
        try:
            if 'gallery_images' not in workbook.sheetnames:
                raise ValueError('没有找到 gallery_images 表，请先扫描图片写入画廊。')
            sheet = workbook['gallery_images']
            columns = self.columns(sheet)
            seen = set()
            entries = []
            for row in sheet.iter_rows(min_row=2, values_only=True):
                filename = str(row[columns['filename'] - 1] or '').strip()
                if not is_gallery_image(filename):
                    continue
                key = gallery_image_key(filename)
                if key in seen:
                    continue
                seen.add(key)
                # 只展示 images 顶层的本地文件；不把任意路径交给图片加载器。
                if '/' in key or key in {'.', '..'}:
                    continue
                image_path = self.root / 'images' / key
                entries.append({'key': key, 'filename': filename, 'path': image_path,
                                'tags': str(row[columns['tags'] - 1] or '')})
            self.entries = entries
        finally:
            workbook.close()

    def save(self, entry, raw):
        raw = validate_tags(raw)
        backup = self.save_many([(dict(entry), raw)])
        entry['tags'] = raw
        return backup

    def save_many(self, changes):
        """一个批次只读写一次工作簿；传入快照，不修改界面内存。"""
        changes = [(entry, validate_tags(raw)) for entry, raw in changes]
        if (self.path.parent / ('~$' + self.path.name)).exists():
            raise PermissionError('Excel 正在打开 data.xlsx。请关闭 Excel 后重试，当前填写内容会保留。')
        before = self.path.read_bytes()
        workbook = openpyxl.load_workbook(self.path)
        temporary = None
        try:
            sheet = workbook['gallery_images']
            columns = self.columns(sheet)
            rows = {}
            for row in range(2, sheet.max_row + 1):
                key = gallery_image_key(sheet.cell(row, columns['filename']).value)
                rows.setdefault(key, []).append(row)
            changed = False
            for entry, raw in changes:
                matches = rows.get(entry['key'], [])
                if len(matches) != 1:
                    raise ValueError('图片记录已删除或出现重复：' + entry['filename'])
                cell = sheet.cell(matches[0], columns['tags'])
                if str(cell.value or '') != entry['tags']:
                    raise ValueError('标签已被其他操作修改，未覆盖：' + entry['filename'])
                if raw != entry['tags']:
                    cell.value = raw or None
                    changed = True
            if not changed:
                return None
            descriptor, name = tempfile.mkstemp(prefix='.gallery-tags-', suffix='.xlsx', dir=self.path.parent)
            os.close(descriptor)
            temporary = Path(name)
            workbook.save(temporary)
            workbook.close()
            if hashlib.sha256(self.path.read_bytes()).digest() != hashlib.sha256(before).digest():
                raise ValueError('保存期间表格发生了变化，请重试。当前填写内容已保留。')
            backup = self.root / '_backups' / ('data-gallery-tags-' + datetime.now().strftime('%Y%m%d-%H%M%S-%f') + '.xlsx')
            backup.parent.mkdir(exist_ok=True)
            shutil.copy2(self.path, backup)
            os.replace(temporary, self.path)
            return backup
        finally:
            workbook.close()
            if temporary is not None and temporary.exists():
                temporary.unlink()


class GalleryEditor(tk.Toplevel):
    PAGE_SIZE = 30

    def __init__(self, parent, root):
        super().__init__(parent)
        self.withdraw()
        try:
            self.store = GalleryStore(root)
        except Exception as error:
            messagebox.showerror('无法打开画廊', str(error), parent=parent)
            self.destroy()
            return
        self.title('看图填标签 · 唯鸡百科')
        self.geometry('1200x820')
        self.minsize(960, 650)
        self.current = None
        self.filtered = list(range(len(self.store.entries)))
        self.page = 0
        self.last_saved = None
        self.persisted = {entry['key']: entry['tags'] for entry in self.store.entries}
        self.pending = {}
        self.saving = False
        self.save_results = queue.Queue()
        self.flush_job = None
        self.poll_job = None
        self.closing = False
        self.on_closed = None
        self.vocabulary_signature = None
        self.thumb_images = {}
        self.preview_image = None
        self.render_job = None
        self.preview_job = None
        self.loading = False
        self.raw = tk.StringVar()
        self.query = tk.StringVar()
        self.only_empty = tk.BooleanVar()
        self.status = tk.StringVar(value='标签只写入 data.xlsx；完成后在主窗口点击“更新所有 data”。')
        existing_categories = {category for entry in self.store.entries
                               for category in parse_gallery_tags(entry['tags'])}
        self.dimensions = (sorted(existing_categories, key=lambda c: (DIMENSIONS.index(c) if c in DIMENSIONS else 99, c))
                           if existing_categories else list(DIMENSIONS))
        self.inherit = {category: tk.BooleanVar(value=category in {'类型', 'LIVE', '分类', '时期', '场次'})
                        for category in self.dimensions}
        self.checks = {}
        self.build_ui()
        self.raw.trace_add('write', self.sync_checks)
        self.protocol('WM_DELETE_WINDOW', self.close)
        self.bind('<Control-s>', lambda event: self.save_key(False))
        self.bind('<Control-Return>', lambda event: self.save_key(True))
        self.bind('<Alt-Left>', lambda event: self.navigate(-1))
        self.bind('<Alt-Right>', lambda event: self.navigate(1))
        self.refresh_tags()
        self.refresh_list()
        if self.filtered:
            self.show_entry(self.filtered[0])
        self.deiconify()
        self.grab_set()
        self.poll_job = self.after(60, self.poll_save)

    def build_ui(self):
        header = ttk.Frame(self, padding=12)
        header.pack(fill='x')
        ttk.Label(header, text='看图填标签', font=('Microsoft YaHei', 15, 'bold')).pack(side='left')
        ttk.Label(header, text='Ctrl+S 保存 · Ctrl+Enter 保存并下一张 · Alt+←/→ 切换图片').pack(side='right')
        ttk.Label(self, textvariable=self.status, padding=12, wraplength=1100).pack(side='bottom', fill='x')
        panes = ttk.Panedwindow(self, orient='horizontal')
        panes.pack(fill='both', expand=True, padx=12)
        left = ttk.Frame(panes, width=310)
        right = ttk.Frame(panes)
        panes.add(left, weight=1)
        panes.add(right, weight=3)
        right.columnconfigure(0, weight=1)
        right.rowconfigure(1, weight=1)
        search = ttk.Entry(left, textvariable=self.query)
        search.pack(fill='x', pady=(0, 5))
        search.bind('<Return>', lambda event: self.filter_list())
        bar = ttk.Frame(left)
        bar.pack(fill='x')
        ttk.Button(bar, text='搜索文件名 / 标签', command=self.filter_list).pack(side='left')
        ttk.Checkbutton(bar, text='仅未填标签', variable=self.only_empty, command=self.filter_list).pack(side='left')
        tree_frame = ttk.Frame(left)
        tree_frame.pack(fill='both', expand=True, pady=8)
        style = ttk.Style(self)
        style.configure('Gallery.Treeview', rowheight=66)
        self.tree = ttk.Treeview(tree_frame, show='tree', selectmode='browse', style='Gallery.Treeview')
        self.tree.column('#0', width=280, minwidth=180)
        scroll = ttk.Scrollbar(tree_frame, command=self.tree.yview)
        self.tree.configure(yscrollcommand=scroll.set)
        scroll.pack(side='right', fill='y')
        self.tree.pack(fill='both', expand=True)
        self.tree.bind('<<TreeviewSelect>>', self.select_entry)
        pager = ttk.Frame(left)
        pager.pack(fill='x', pady=(0, 8))
        ttk.Button(pager, text='上一页', command=lambda: self.change_page(-1)).pack(side='left')
        self.page_label = ttk.Label(pager)
        self.page_label.pack(side='left', expand=True)
        ttk.Button(pager, text='下一页', command=lambda: self.change_page(1)).pack(side='right')

        self.filename_label = ttk.Label(right, text='没有图片，请先扫描图片写入画廊。', wraplength=780)
        self.filename_label.grid(row=0, column=0, sticky='ew', padx=(12, 0))
        self.preview = tk.Canvas(right, bg='#171717', highlightthickness=0, height=300)
        self.preview.grid(row=1, column=0, sticky='nsew', padx=(12, 0), pady=6)
        self.preview.bind('<Configure>', self.schedule_preview)
        self.preview.bind('<Double-Button-1>', self.enlarge)
        ttk.Label(right, text='双击图片放大。人物＝确认出镜；关联人物＝访谈或活动相关。').grid(row=2, column=0, sticky='w', padx=12)
        self.notebook = ttk.Notebook(right, height=140)
        self.notebook.grid(row=3, column=0, sticky='ew', padx=(12, 0), pady=6)
        custom = ttk.Frame(right)
        custom.grid(row=4, column=0, sticky='ew', padx=(12, 0))
        self.category = ttk.Combobox(custom, values=self.dimensions, width=12)
        self.category.set('分类' if '分类' in self.dimensions else '类型' if '类型' in self.dimensions else self.dimensions[0])
        self.category.pack(side='left')
        self.new_tag = ttk.Entry(custom)
        self.new_tag.pack(side='left', fill='x', expand=True, padx=6)
        self.new_tag.bind('<Return>', lambda event: self.add_tag())
        ttk.Button(custom, text='添加标签', command=self.add_tag).pack(side='right')
        ttk.Label(right, text='完整标签（可直接编辑，清空后保存即移除标签）').grid(row=5, column=0, sticky='w', padx=12, pady=(6, 0))
        ttk.Entry(right, textvariable=self.raw).grid(row=6, column=0, sticky='ew', padx=(12, 0), pady=4)
        self.inherit_frame = ttk.Frame(right)
        self.inherit_frame.grid(row=7, column=0, sticky='ew', padx=(12, 0))
        self.refresh_inherit_controls()
        buttons = ttk.Frame(right)
        buttons.grid(row=8, column=0, sticky='ew', padx=(12, 0), pady=10)
        for text, command in [('上一张', lambda: self.navigate(-1)), ('下一张', lambda: self.navigate(1)),
                              ('沿用上一张', self.inherit_tags), ('保存', self.save),
                              ('保存并下一张', lambda: self.save(next_image=True))]:
            ttk.Button(buttons, text=text, command=command).pack(side='left', padx=(0, 6))

    def refresh_tags(self):
        vocabulary = {}
        for entry in self.store.entries:
            for category, values in parse_gallery_tags(entry['tags']).items():
                vocabulary.setdefault(category, set()).update(values)
        for category, values in parse_gallery_tags(self.raw.get()).items():
            vocabulary.setdefault(category, set()).update(values)
        if not vocabulary:
            vocabulary = {key: set(values) for key, values in DEFAULT_TAGS.items()}
        signature = tuple(sorted((key, tuple(sorted(values))) for key, values in vocabulary.items()))
        if signature == self.vocabulary_signature:
            self.sync_checks()
            return
        self.vocabulary_signature = signature
        for category in vocabulary:
            if category not in self.inherit:
                self.dimensions.append(category)
                self.inherit[category] = tk.BooleanVar(value=False)
        self.category.configure(values=self.dimensions)
        self.refresh_inherit_controls()
        for tab in self.notebook.tabs():
            self.notebook.nametowidget(tab).destroy()
        self.checks = {}
        categories = sorted(vocabulary, key=lambda c: (DIMENSIONS.index(c) if c in DIMENSIONS else 99, c))
        for category in categories:
            frame = ttk.Frame(self.notebook)
            self.notebook.add(frame, text=category)
            canvas = tk.Canvas(frame, highlightthickness=0, height=140)
            scroll = ttk.Scrollbar(frame, command=canvas.yview)
            canvas.configure(yscrollcommand=scroll.set)
            scroll.pack(side='right', fill='y')
            canvas.pack(side='left', fill='both', expand=True)
            content = ttk.Frame(canvas)
            canvas.create_window((0, 0), window=content, anchor='nw')
            content.bind('<Configure>', lambda event, c=canvas: c.configure(scrollregion=c.bbox('all')))
            canvas.bind('<MouseWheel>', lambda event, c=canvas: c.yview_scroll(-int(event.delta / 120), 'units'))
            for index, value in enumerate(sorted(vocabulary[category])):
                selected = tk.BooleanVar()
                self.checks[(category, value)] = selected
                ttk.Checkbutton(content, text=value, variable=selected,
                                command=lambda c=category, v=value: self.toggle_tag(c, v)).grid(
                                    row=index // 3, column=index % 3, sticky='w', padx=8, pady=3)
        self.sync_checks()

    def refresh_inherit_controls(self):
        for widget in self.inherit_frame.winfo_children():
            widget.destroy()
        ttk.Label(self.inherit_frame, text='沿用范围：').grid(row=0, column=0, sticky='w')
        for index, category in enumerate(self.dimensions):
            ttk.Checkbutton(self.inherit_frame, text=category, variable=self.inherit[category]).grid(
                row=index // 5, column=index % 5 + 1, sticky='w')

    def sync_checks(self, *_):
        groups = parse_gallery_tags(self.raw.get())
        for (category, value), variable in self.checks.items():
            variable.set(value in groups.get(category, []))

    def toggle_tag(self, category, value):
        groups = parse_gallery_tags(self.raw.get())
        values = groups.setdefault(category, [])
        if value in values:
            values.remove(value)
            if not values:
                groups.pop(category)
        else:
            values.append(value)
        self.raw.set(format_tags(groups))

    def add_tag(self):
        category, value = self.category.get().strip(), self.new_tag.get().strip()
        if not category or not value or any(c in category + value for c in ',:，：\n'):
            messagebox.showerror('标签格式', '请填写分类和标签，名称不能包含逗号、冒号或换行。', parent=self)
            return
        groups = parse_gallery_tags(self.raw.get())
        if value not in groups.setdefault(category, []):
            groups[category].append(value)
        self.raw.set(format_tags(groups))
        self.new_tag.delete(0, 'end')
        self.refresh_tags()

    def filter_list(self):
        if not self.leave_current():
            return
        query = self.query.get().strip().casefold()
        self.filtered = [i for i, entry in enumerate(self.store.entries)
                         if (not self.only_empty.get() or not entry['tags'].strip())
                         and query in (entry['filename'] + ' ' + entry['tags']).casefold()]
        self.page = 0
        self.refresh_list()
        self.show_entry(self.filtered[0] if self.filtered else None)

    def refresh_list(self):
        if self.render_job:
            self.after_cancel(self.render_job)
        self.loading = True
        self.tree.delete(*self.tree.get_children())
        self.thumb_images.clear()
        visible = self.filtered[self.page * self.PAGE_SIZE:(self.page + 1) * self.PAGE_SIZE]
        for i in visible:
            entry = self.store.entries[i]
            self.tree.insert('', 'end', iid=str(i), text=('✓ ' if entry['tags'].strip() else '○ ') + Path(entry['filename']).name)
        self.page_label.config(text=f'{self.page + 1}/{max(1, (len(self.filtered) + self.PAGE_SIZE - 1) // self.PAGE_SIZE)} · {len(self.filtered)} 张')
        self.loading = False
        self.render_thumbnails(visible)

    def render_thumbnails(self, indices):
        self.render_job = None
        if not indices:
            return
        index, *remaining = indices
        try:
            with Image.open(self.store.entries[index]['path']) as source:
                thumb = ImageOps.exif_transpose(source)
                thumb.thumbnail((78, 60))
                photo = ImageTk.PhotoImage(thumb, master=self)
            self.thumb_images[index] = photo
            if self.tree.exists(str(index)):
                self.tree.item(str(index), image=photo)
        except (OSError, ValueError):
            pass
        self.render_job = self.after(5, lambda: self.render_thumbnails(remaining))

    def select_entry(self, _event=None):
        selection = self.tree.selection()
        if self.loading or not selection or int(selection[0]) == self.current:
            return
        index = int(selection[0])
        if self.leave_current():
            self.show_entry(index)
        elif self.current is not None and self.tree.exists(str(self.current)):
            self.tree.selection_set(str(self.current))

    def show_entry(self, index):
        self.current = index
        entry = self.store.entries[index] if index is not None else None
        self.raw.set(entry['tags'] if entry else '')
        self.filename_label.config(text=entry['filename'] if entry else '没有符合条件的图片。')
        if entry and self.tree.exists(str(index)):
            self.tree.selection_set(str(index))
            self.tree.see(str(index))
        self.draw_preview()

    def schedule_preview(self, _event=None):
        if self.preview_job:
            self.after_cancel(self.preview_job)
        self.preview_job = self.after(120, self.draw_preview)

    def draw_preview(self):
        self.preview_job = None
        self.preview.delete('all')
        self.preview_image = None
        if self.current is None:
            return
        width, height = max(100, self.preview.winfo_width()), max(100, self.preview.winfo_height())
        try:
            with Image.open(self.store.entries[self.current]['path']) as source:
                picture = ImageOps.exif_transpose(source)
                picture.thumbnail((width - 12, height - 12))
                self.preview_image = ImageTk.PhotoImage(picture, master=self)
            self.preview.create_image(width // 2, height // 2, image=self.preview_image)
        except (OSError, ValueError):
            self.preview.create_text(width // 2, height // 2, text='图片不存在或无法读取，仍可编辑标签。', fill='white')

    def enlarge(self, _event=None):
        if self.current is None:
            return
        try:
            with Image.open(self.store.entries[self.current]['path']) as source:
                picture = ImageOps.exif_transpose(source)
                picture.thumbnail((self.winfo_screenwidth() - 120, self.winfo_screenheight() - 160))
                window = tk.Toplevel(self)
                window.title(self.store.entries[self.current]['filename'])
                photo = ImageTk.PhotoImage(picture, master=window)
                label = tk.Label(window, image=photo, bg='#171717')
                label.image = photo
                label.pack()
                window.bind('<Escape>', lambda event: window.destroy())
        except (OSError, ValueError) as error:
            messagebox.showerror('无法打开图片', str(error), parent=self)

    def leave_current(self):
        if self.current is None or self.raw.get() == self.store.entries[self.current]['tags']:
            return True
        decision = messagebox.askyesnocancel('尚未保存', '是否保存当前图片的标签？\n选择“否”放弃本张修改，选择“取消”继续编辑。', parent=self)
        return self.save() if decision else decision is False

    def navigate(self, step):
        if self.current not in self.filtered or not self.leave_current():
            return
        position = self.filtered.index(self.current) + step
        if not 0 <= position < len(self.filtered):
            self.status.set('已经是第一张或最后一张。')
            return
        page = position // self.PAGE_SIZE
        if page != self.page:
            self.page = page
            self.refresh_list()
        self.show_entry(self.filtered[position])

    def change_page(self, step):
        page = self.page + step
        if not 0 <= page < max(1, (len(self.filtered) + self.PAGE_SIZE - 1) // self.PAGE_SIZE) or not self.leave_current():
            return
        self.page = page
        self.refresh_list()
        self.show_entry(self.filtered[page * self.PAGE_SIZE])

    def inherit_tags(self):
        if self.last_saved is None:
            self.status.set('请先保存一张图片，然后在下一张中沿用它的标签。')
            return
        groups = parse_gallery_tags(self.raw.get())
        previous = parse_gallery_tags(self.last_saved)
        for category, enabled in self.inherit.items():
            if enabled.get():
                groups.pop(category, None)
                if category in previous:
                    groups[category] = previous[category][:]
        self.raw.set(format_tags(groups))
        self.status.set('已沿用上一张保存的所选分类；其他分类保持原样，尚未保存。')

    def save_key(self, next_image):
        self.save(next_image)
        return 'break'

    def save(self, next_image=False):
        if self.current is None:
            if self.pending:
                self.flush_pending()
                return True
            return False
        entry = self.store.entries[self.current]
        try:
            raw = validate_tags(self.raw.get())
        except Exception as error:
            messagebox.showerror('未能保存，填写内容已保留', str(error), parent=self)
            return False
        if raw != entry['tags']:
            self.pending[entry['key']] = raw
            entry['tags'] = raw
        self.raw.set(entry['tags'])
        self.last_saved = entry['tags']
        self.status.set(f'待写入 {len(self.pending)} 张，正在后台保存；可继续填图。' if self.pending
                        else '标签已保存，内容未变化。')
        self.refresh_tags()
        if self.tree.exists(str(self.current)):
            self.tree.item(str(self.current), text=('✓ ' if entry['tags'] else '○ ') + Path(entry['filename']).name)
        if next_image:
            self.navigate(1)
        if self.pending and self.flush_job is None and not self.saving:
            self.flush_job = self.after(400, self.flush_pending)
        return True

    def flush_pending(self):
        if self.flush_job:
            self.after_cancel(self.flush_job)
            self.flush_job = None
        if self.saving or not self.pending:
            return
        snapshot = dict(self.pending)
        changes = [({'key': key, 'filename': key, 'tags': self.persisted[key]}, raw)
                   for key, raw in snapshot.items()]
        self.saving = True

        def write():
            try:
                self.store.save_many(changes)
                self.save_results.put((snapshot, None))
            except Exception as error:
                self.save_results.put((snapshot, error))
        threading.Thread(target=write, daemon=True).start()

    def poll_save(self):
        self.poll_job = None
        try:
            snapshot, error = self.save_results.get_nowait()
        except queue.Empty:
            pass
        else:
            self.saving = False
            if error:
                self.closing = False
                self.on_closed = None
                self.deiconify()
                self.grab_set()
                self.status.set(f'写入失败，{len(self.pending)} 张修改仍在待保存队列。处理后点击“保存”重试。')
                messagebox.showerror('未写入 Excel，修改已保留', str(error), parent=self)
            else:
                for key, raw in snapshot.items():
                    self.persisted[key] = raw
                    if self.pending.get(key) == raw:
                        self.pending.pop(key)
                if self.pending:
                    self.flush_pending()
                else:
                    self.status.set('全部标签已写入 data.xlsx，原表已备份。关闭面板后可更新所有 data。')
                    if self.closing:
                        self.finish_close()
                        return
        self.poll_job = self.after(60, self.poll_save)

    def close(self, on_closed=None):
        if self.closing:
            if on_closed:
                self.on_closed = on_closed
            return
        if not self.leave_current():
            return
        self.on_closed = on_closed
        if self.pending or self.saving:
            self.closing = True
            self.status.set('正在完成最后的保存，请稍候…')
            self.withdraw()
            self.flush_pending()
        else:
            self.finish_close()

    def finish_close(self):
        for job in (self.render_job, self.preview_job, self.flush_job, self.poll_job):
            if job:
                self.after_cancel(job)
        callback = self.on_closed
        self.grab_release()
        self.destroy()
        if callback:
            callback()


def smoke_test():
    """在临时工作簿上验证真实 Tk 面板，供源码和打包后的 EXE 共用。"""
    import time
    with tempfile.TemporaryDirectory() as directory:
        root = Path(directory)
        (root / '_data').mkdir()
        (root / 'images').mkdir()
        for name in ('a.webp', 'b.webp'):
            Image.new('RGB', (120, 80), 'red').save(root / 'images' / name)
        workbook = openpyxl.Workbook()
        sheet = workbook.active
        sheet.title = 'gallery_images'
        sheet.append(['filename', 'tags'])
        sheet.append(['../images/a.webp', '类型:访谈, LIVE:0th'])
        sheet.append(['../images/b.webp', '人物:高尾奏音'])
        workbook.save(root / '_data/data.xlsx')
        workbook.close()
        parent = tk.Tk()
        parent.withdraw()
        editor = GalleryEditor(parent, root)
        editor.withdraw()
        def until(condition):
            deadline = time.monotonic() + 10
            while not condition():
                parent.update()
                if time.monotonic() > deadline:
                    raise AssertionError('background save did not finish')
                time.sleep(0.01)
        try:
            for _ in range(15):
                parent.update()
                time.sleep(0.01)
            assert editor.current == 0
            assert editor.preview_image is not None
            assert len(editor.thumb_images) == 2
            editor.toggle_tag('人物', '佐佐木李子')
            assert editor.save(next_image=True)
            assert editor.current == 1
            editor.inherit_tags()
            groups = parse_gallery_tags(editor.raw.get())
            assert groups['类型'] == ['访谈'] and groups['LIVE'] == ['0th']
            assert groups['人物'] == ['高尾奏音']
            assert editor.save()
            # 连续两张合并成一个 Excel 写入批次，点击保存时尚未进行磁盘写入。
            assert len(editor.pending) == 2 and not editor.saving
            calls = []
            real_save_many = editor.store.save_many
            def record(changes):
                calls.append(changes)
                return real_save_many(changes)
            editor.store.save_many = record
            editor.flush_pending()
            until(lambda: not editor.pending and not editor.saving)
            assert len(calls) == 1 and len(calls[0]) == 2

            # 写入期间继续改同一张，旧批次完成不能抹掉后续修改。
            entered, release = threading.Event(), threading.Event()
            def slow_write(changes):
                entered.set()
                assert release.wait(5)
                return real_save_many(changes)
            editor.store.save_many = slow_write
            editor.raw.set('类型:合照')
            assert editor.save()
            editor.flush_pending()
            until(entered.is_set)
            editor.raw.set('类型:访谈, 人物:高尾奏音')
            assert editor.save()
            assert editor.saving
            release.set()
            until(lambda: not editor.pending and not editor.saving)
            editor.store.save_many = real_save_many

            # 写入失败时队列和草稿保留，重试成功后才算真正保存。
            original_showerror = messagebox.showerror
            errors = []
            messagebox.showerror = lambda *args, **kwargs: errors.append(args)
            try:
                def fail_write(changes):
                    raise PermissionError('simulated Excel lock')
                editor.store.save_many = fail_write
                editor.raw.set('类型:日常')
                assert editor.save()
                editor.flush_pending()
                until(lambda: bool(errors))
                assert editor.pending and editor.raw.get() == '类型:日常'
                editor.store.save_many = real_save_many
                editor.raw.set('类型:访谈, 人物:高尾奏音')
                assert editor.save()
                editor.flush_pending()
                until(lambda: not editor.pending and not editor.saving)
            finally:
                messagebox.showerror = original_showerror
                editor.store.save_many = real_save_many
            editor.query.set('b.webp')
            editor.filter_list()
            assert editor.filtered == [1]
            editor.query.set('no-such-image')
            editor.filter_list()
            assert editor.current is None
            editor.query.set('')
            editor.only_empty.set(True)
            editor.filter_list()
            assert not editor.filtered
            saved = GalleryStore(root)
            assert '人物:佐佐木李子' in saved.entries[0]['tags']
            assert '类型:访谈' in saved.entries[1]['tags']
            # 关闭时等待最后一批写完，回调只在持久化成功之后执行。
            editor.only_empty.set(False)
            editor.filter_list()
            editor.raw.set('类型:合照')
            assert editor.save()
            closed = []
            editor.close(on_closed=lambda: closed.append(True))
            until(lambda: bool(closed))
            assert GalleryStore(root).entries[0]['tags'] == '类型:合照'
        finally:
            if editor.winfo_exists():
                editor.close()
            parent.destroy()
