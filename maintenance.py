#!/usr/bin/env python3
"""wijipedia 本地维护工具：体检、完整更新与变更报告。"""

from __future__ import annotations

import argparse
import hashlib
import json
import os
import re
import shutil
import subprocess
import sys
from dataclasses import dataclass, field
from datetime import date, datetime
from pathlib import Path
from typing import Callable

import openpyxl


ROOT = Path(sys.executable).resolve().parent if getattr(sys, "frozen", False) else Path(__file__).resolve().parent
SEARCH_INDEX_PREFIX = "window.WIJIPEDIA_SEARCH_INDEX = "
GENERATED_FILES = (
    "announcements.js", "something-new.js", "search-index.js", "search-bodies.js",
    "songs/data.js", "songs/lyrics-atlas-data.js", "live/data.js", "timeline/data.js",
    "gallery/data.js", "interview/data.js", "discography/data.js",
)
REQUIRED_SHEETS = (
    "announcements", "something_new", "songs", "lives", "timeline", "gallery_images",
    "interview", "discography_releases",
)
REQUIRED_FIELDS = {
    "announcements": ("date", "msg"),
    "something_new": ("content_type", "content_ref", "update_date"),
    "songs": ("song_name", "release_date"),
    "lives": ("live_name", "live_date"),
    "timeline": ("date", "title"),
    "gallery_images": ("filename",),
    "interview": ("title", "md_path"),
    "discography_releases": ("release_id", "title", "release_date"),
    "setlist": ("live_name", "track_title"),
    "song_comments": ("song_name", "comment_text"),
    "song_live_history": ("song_name", "live_name"),
    "discography_editions": ("release_id", "edition_name"),
    "discography_contents": ("release_id", "edition_name"),
    "discography_bonus_contents": ("release_id", "bonus_id"),
}
DATE_FIELDS = {
    "announcements": ("date",), "something_new": ("update_date",), "songs": ("release_date",),
    "lives": ("live_date",), "timeline": ("date",), "gallery_images": ("date",),
    "interview": ("date",), "discography_releases": ("release_date",),
    "song_comments": ("comment_date",), "song_live_history": ("live_date",),
}
UNIQUE_FIELDS = {
    "songs": "song_name", "lives": "live_name",
    "interview": "title", "discography_releases": "release_id",
}


@dataclass
class Issue:
    severity: str
    location: str
    message: str


@dataclass
class CheckReport:
    issues: list[Issue] = field(default_factory=list)
    counts: dict[str, int] = field(default_factory=dict)

    @property
    def errors(self) -> list[Issue]:
        return [item for item in self.issues if item.severity == "ERROR"]

    @property
    def warnings(self) -> list[Issue]:
        return [item for item in self.issues if item.severity == "WARN"]


def _text(value) -> str:
    return "" if value is None else str(value).strip()


def _rows(sheet) -> list[tuple[int, dict[str, object]]]:
    iterator = sheet.iter_rows(values_only=True)
    try:
        raw_headers = next(iterator)
    except StopIteration:
        return []
    headers = [_text(value) for value in raw_headers]
    result = []
    for row_number, values in enumerate(iterator, start=2):
        row = {headers[index]: values[index] for index in range(min(len(headers), len(values))) if headers[index]}
        if any(_text(value) for value in row.values()):
            result.append((row_number, row))
    return result


def _valid_date(value) -> bool:
    if isinstance(value, (datetime, date)):
        return True
    text = _text(value)
    if not text:
        return True
    match = re.fullmatch(r"(\d{4})[./-](\d{1,2})[./-](\d{1,2})(?:\s+.*)?", text)
    if not match:
        return False
    try:
        date(int(match.group(1)), int(match.group(2)), int(match.group(3)))
        return True
    except ValueError:
        return False


def _asset_candidates(root: Path, value: str) -> list[Path]:
    cleaned = value.replace("\\", "/").strip()
    while cleaned.startswith("../"):
        cleaned = cleaned[3:]
    candidates = [root / cleaned]
    if not cleaned.startswith(("images/", "icons/", "_data/")):
        candidates.extend((root / "images" / cleaned, root / "_data" / cleaned))
    return candidates


def check_project(root: Path | str = ROOT, log_func: Callable[[str], None] | None = None) -> CheckReport:
    root = Path(root)
    report = CheckReport()
    xlsx = root / "_data" / "data.xlsx"

    def add(severity: str, location: str, message: str):
        report.issues.append(Issue(severity, location, message))

    if not xlsx.exists():
        add("ERROR", "_data/data.xlsx", "数据源不存在")
        return report

    try:
        workbook = openpyxl.load_workbook(xlsx, data_only=True, read_only=True)
    except Exception as exc:
        add("ERROR", "_data/data.xlsx", f"无法读取：{exc}")
        return report

    try:
        for sheet_name in REQUIRED_SHEETS:
            if sheet_name not in workbook.sheetnames:
                add("ERROR", sheet_name, "缺少必要工作表")

        cached_rows: dict[str, list[tuple[int, dict[str, object]]]] = {}
        for sheet_name in workbook.sheetnames:
            rows = _rows(workbook[sheet_name])
            cached_rows[sheet_name] = rows
            report.counts[sheet_name] = len(rows)
            required = REQUIRED_FIELDS.get(sheet_name, ())
            seen: dict[str, int] = {}
            unique_field = UNIQUE_FIELDS.get(sheet_name)
            for row_number, row in rows:
                location = f"{sheet_name}!{row_number}"
                for field_name in required:
                    if not _text(row.get(field_name)):
                        add("ERROR", location, f"必填字段 {field_name} 为空")
                for field_name in DATE_FIELDS.get(sheet_name, ()):
                    if not _valid_date(row.get(field_name)):
                        add("ERROR", location, f"日期 {field_name} 格式无效：{_text(row.get(field_name))}")
                if unique_field:
                    value = _text(row.get(unique_field))
                    if value:
                        if value in seen:
                            add("ERROR", location, f"{unique_field} 与第 {seen[value]} 行重复：{value}")
                        else:
                            seen[value] = row_number

        songs = {_text(row.get("song_name")) for _, row in cached_rows.get("songs", [])}
        lives = {_text(row.get("live_name")) for _, row in cached_rows.get("lives", [])}
        releases = {_text(row.get("release_id")) for _, row in cached_rows.get("discography_releases", [])}
        references = (
            ("song_comments", "song_name", songs, "歌曲"),
            ("song_live_history", "song_name", songs, "歌曲"),
            ("setlist", "live_name", lives, "演出"),
            ("backstage", "live_name", lives, "演出"),
            ("discography_editions", "release_id", releases, "发行作品"),
            ("discography_contents", "release_id", releases, "发行作品"),
            ("discography_bonus_contents", "release_id", releases, "发行作品"),
        )
        for sheet_name, field_name, known, label in references:
            for row_number, row in cached_rows.get(sheet_name, []):
                value = _text(row.get(field_name))
                if value and value not in known:
                    add("ERROR", f"{sheet_name}!{row_number}", f"引用的{label}不存在：{value}")

        asset_fields = {
            "songs": ("cover",), "lives": ("poster", "kv"), "gallery_images": ("filename",),
            "interview": ("poster",), "discography_releases": ("cover",),
            "discography_editions": ("cover",), "backstage": ("photo",),
        }
        for sheet_name, fields in asset_fields.items():
            for row_number, row in cached_rows.get(sheet_name, []):
                for field_name in fields:
                    value = _text(row.get(field_name))
                    if value and not value.startswith(("http://", "https://")):
                        if not any(candidate.is_file() for candidate in _asset_candidates(root, value)):
                            add("WARN", f"{sheet_name}!{row_number}", f"{field_name} 指向的文件不存在：{value}")

        for sheet_name, field_name in (("interview", "md_path"), ("setlist", "mc_file")):
            for row_number, row in cached_rows.get(sheet_name, []):
                value = _text(row.get(field_name))
                if value and not (root / "_data" / value.lstrip("/\\")).is_file():
                    add("ERROR", f"{sheet_name}!{row_number}", f"Markdown 文件不存在：{value}")

        for row_number, row in cached_rows.get("interview", []):
            translated = _text(row.get("if_translated")).lower()
            if translated and translated not in ("yes", "no"):
                add("WARN", f"interview!{row_number}", f"if_translated 应为 yes/no：{translated}")
    finally:
        workbook.close()

    report.issues.sort(key=lambda item: ({"ERROR": 0, "WARN": 1}.get(item.severity, 2), item.location))
    if log_func:
        for line in format_check_report(report).splitlines():
            log_func(line)
    return report


def format_check_report(report: CheckReport) -> str:
    lines = [f"项目体检：{len(report.errors)} 个错误 · {len(report.warnings)} 个提醒"]
    for issue in report.issues:
        lines.append(f"[{issue.severity}] {issue.location} · {issue.message}")
    if not report.issues:
        lines.append("[OK] 未发现数据结构、引用或资源路径问题")
    total = sum(report.counts.get(name, 0) for name in REQUIRED_SHEETS)
    lines.append(f"[INFO] 已检查 {len(report.counts)} 个工作表 · 核心条目 {total} 条")
    return "\n".join(lines)


def _read_search_index_text(text: str) -> list[dict]:
    text = text.strip()
    if not text.startswith(SEARCH_INDEX_PREFIX):
        return []
    payload = text[len(SEARCH_INDEX_PREFIX):].rstrip(";\r\n ")
    try:
        data = json.loads(payload)
        return data if isinstance(data, list) else []
    except json.JSONDecodeError:
        return []


def read_search_index(path: Path) -> list[dict]:
    return _read_search_index_text(path.read_text(encoding="utf-8")) if path.exists() else []


def _file_hash(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest() if path.is_file() else ""


def capture_generated_state(root: Path | str = ROOT) -> dict:
    root = Path(root)
    return {
        "index": read_search_index(root / "search-index.js"),
        "hashes": {name: _file_hash(root / name) for name in GENERATED_FILES},
    }


def compare_states(before: dict, after: dict) -> dict:
    old = {item.get("key"): item for item in before.get("index", []) if item.get("key")}
    new = {item.get("key"): item for item in after.get("index", []) if item.get("key")}
    added = [new[key] for key in sorted(new.keys() - old.keys())]
    removed = [old[key] for key in sorted(old.keys() - new.keys())]
    modified = []
    visible_fields = ("title", "subtitle", "date", "image", "url")
    for key in sorted(old.keys() & new.keys()):
        changed = [field for field in visible_fields if old[key].get(field) != new[key].get(field)]
        if not changed and old[key].get("terms") != new[key].get("terms"):
            changed = ["检索内容"]
        if changed:
            modified.append({"item": new[key], "fields": changed})
    changed_files = [
        name for name in GENERATED_FILES
        if before.get("hashes", {}).get(name) != after.get("hashes", {}).get(name)
    ]
    return {"added": added, "removed": removed, "modified": modified, "changed_files": changed_files}


def format_change_report(changes: dict) -> str:
    lines = [
        "内容变更报告",
        f"新增 {len(changes['added'])} · 删除 {len(changes['removed'])} · 修改 {len(changes['modified'])} · 生成文件 {len(changes['changed_files'])}",
    ]
    for label, items, prefix in (("新增", changes["added"], "+"), ("删除", changes["removed"], "-")):
        if items:
            lines.append(f"\n{label}：")
            for item in items:
                lines.append(f"  {prefix} [{item.get('type', 'file')}] {item.get('title') or item.get('subtitle') or item.get('key')} · {item.get('date') or '无日期'}")
    if changes["modified"]:
        lines.append("\n修改：")
        for change in changes["modified"]:
            item = change["item"]
            lines.append(f"  ~ [{item.get('type', 'file')}] {item.get('title') or item.get('key')} · {', '.join(change['fields'])}")
    if changes["changed_files"]:
        lines.append("\n已更新生成文件：")
        lines.extend(f"  · {name}" for name in changes["changed_files"])
    if not any((changes["added"], changes["removed"], changes["modified"], changes["changed_files"])):
        lines.append("\n[OK] 没有检测到内容变化")
    return "\n".join(lines)


def _configure_generator(root: Path):
    import generate_all
    generate_all.ROOT = str(root)
    generate_all.DATA_DIR = str(root / "_data")
    generate_all.DOCS_DIR = str(root / "_data" / "_docs")
    generate_all.XLSX_PATH = str(root / "_data" / "data.xlsx")
    for key, current in list(generate_all.OUTPUTS.items()):
        relative = Path(current).name if key in ("announcements", "something_new", "search_index", "search_bodies") else None
        if relative:
            generate_all.OUTPUTS[key] = str(root / relative)
    generate_all.OUTPUTS.update({
        "songs": str(root / "songs" / "data.js"), "lyrics_atlas": str(root / "songs" / "lyrics-atlas-data.js"),
        "live": str(root / "live" / "data.js"), "timeline": str(root / "timeline" / "data.js"),
        "gallery": str(root / "gallery" / "data.js"), "interview": str(root / "interview" / "data.js"),
        "discography": str(root / "discography" / "data.js"),
    })
    return generate_all


def generate_all_content(root: Path | str = ROOT, log_func: Callable[[str], None] = print) -> dict[str, str]:
    root = Path(root)
    generator = _configure_generator(root)
    workbook = openpyxl.load_workbook(root / "_data" / "data.xlsx", data_only=True)
    results: dict[str, str] = {}
    try:
        steps = (
            ("announcements", "公告", lambda: f"{generator.generate_announcements(workbook)} 条"),
            ("something_new", "首页精选", lambda: f"{generator.generate_something_new(workbook)} 条"),
            ("songs", "歌曲", lambda: f"{generator.generate_songs(workbook)} 首"),
            ("lives", "演唱会", lambda: f"{generator.generate_live(workbook)[0]} 场"),
            ("timeline", "时间轴", lambda: f"{generator.generate_timeline(workbook)} 条"),
            ("gallery_images", "画廊", lambda: f"{generator.generate_gallery(workbook)} 张"),
            ("interview", "访谈", lambda: f"{generator.generate_interview(workbook)} 篇"),
            ("discography_releases", "唱片目录", lambda: f"{generator.generate_discography(workbook)} 张"),
        )
        for sheet_name, label, operation in steps:
            if sheet_name not in workbook.sheetnames:
                continue
            results[label] = operation()
            log_func(f"[OK] {label}：{results[label]}")
        search_count, body_count = generator.generate_search_indexes(workbook)
        results["全局搜索"] = f"{search_count} 条索引 · {body_count} 条正文"
        log_func(f"[OK] 全局搜索：{results['全局搜索']}")
        version = generator.inject_version()
        results["缓存版本"] = version
        log_func(f"[OK] 缓存版本：v={version}")
    finally:
        workbook.close()
    return results


def create_backup(root: Path | str = ROOT) -> Path:
    root = Path(root)
    stamp = datetime.now().strftime("%Y-%m-%d_%H%M%S")
    destination = root / "_backups" / stamp
    destination.mkdir(parents=True, exist_ok=False)
    shutil.copy2(root / "_data" / "data.xlsx", destination / "data.xlsx")
    return destination


def write_report(root: Path | str, text: str, prefix: str = "maintenance") -> Path:
    root = Path(root)
    folder = root / "_reports"
    folder.mkdir(parents=True, exist_ok=True)
    stamp = datetime.now().strftime("%Y-%m-%d_%H%M%S")
    path = folder / f"{prefix}-{stamp}.txt"
    path.write_text(text + "\n", encoding="utf-8")
    (folder / f"latest-{prefix}.txt").write_text(text + "\n", encoding="utf-8")
    return path


def run_full_update(root: Path | str = ROOT, log_func: Callable[[str], None] = print) -> dict:
    root = Path(root)
    precheck = check_project(root)
    for line in format_check_report(precheck).splitlines():
        log_func(line)
    if precheck.errors:
        raise RuntimeError(f"项目体检发现 {len(precheck.errors)} 个错误，已停止更新。")

    backup = create_backup(root)
    log_func(f"[OK] 已备份数据源：{backup.relative_to(root)}")
    before = capture_generated_state(root)
    generated = generate_all_content(root, log_func)
    after = capture_generated_state(root)
    changes = compare_states(before, after)
    change_text = format_change_report(changes)
    for line in change_text.splitlines():
        log_func(line)

    postcheck = check_project(root)
    if postcheck.errors:
        raise RuntimeError(f"生成后体检发现 {len(postcheck.errors)} 个错误。")

    import build_site
    build_site.ROOT = root
    build_site.DIST = root / "dist"
    file_count, byte_count = build_site.build()
    log_func(f"[OK] 发布构建：{file_count} 个文件 · {byte_count / 1024 / 1024:.2f} MB")

    full_report = "\n\n".join((format_check_report(postcheck), change_text, f"发布构建：{file_count} 个文件 · {byte_count / 1024 / 1024:.2f} MB"))
    report_path = write_report(root, full_report)
    log_func(f"[OK] 维护报告：{report_path.relative_to(root)}")
    return {"check": postcheck, "changes": changes, "generated": generated, "backup": backup, "report": report_path, "build": (file_count, byte_count)}


def git_change_report(root: Path | str = ROOT) -> str:
    root = Path(root)
    current = capture_generated_state(root)
    before_text = ""
    try:
        completed = subprocess.run(
            ("git", "show", "HEAD:search-index.js"), cwd=root, check=False,
            stdout=subprocess.PIPE, stderr=subprocess.DEVNULL, text=True, encoding="utf-8", errors="replace",
            creationflags=subprocess.CREATE_NO_WINDOW if os.name == "nt" else 0,
        )
        if completed.returncode == 0:
            before_text = completed.stdout
    except OSError:
        pass
    before = {"index": _read_search_index_text(before_text), "hashes": {}}
    changes = compare_states(before, current)
    changes["changed_files"] = []
    try:
        completed = subprocess.run(
            ("git", "diff", "--name-only", "HEAD", "--", *GENERATED_FILES), cwd=root, check=False,
            stdout=subprocess.PIPE, stderr=subprocess.DEVNULL, text=True, encoding="utf-8", errors="replace",
            creationflags=subprocess.CREATE_NO_WINDOW if os.name == "nt" else 0,
        )
        if completed.returncode == 0:
            changes["changed_files"] = [line for line in completed.stdout.splitlines() if line]
    except OSError:
        pass
    return format_change_report(changes)


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("command", choices=("check", "update", "report"))
    parser.add_argument("--root", type=Path, default=ROOT)
    args = parser.parse_args(argv)
    if args.command == "check":
        report = check_project(args.root)
        print(format_check_report(report))
        return 1 if report.errors else 0
    if args.command == "report":
        text = git_change_report(args.root)
        print(text)
        print(f"\n报告已保存：{write_report(args.root, text, 'changes')}")
        return 0
    run_full_update(args.root)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
