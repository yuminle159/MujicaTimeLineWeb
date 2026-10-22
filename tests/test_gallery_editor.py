import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

import openpyxl
from gallery_editor import GalleryStore


class GalleryStoreTests(unittest.TestCase):
    def setUp(self):
        self.directory = tempfile.TemporaryDirectory()
        self.addCleanup(self.directory.cleanup)
        self.root = Path(self.directory.name)
        (self.root / '_data').mkdir()
        (self.root / 'images').mkdir()
        self.path = self.root / '_data/data.xlsx'
        workbook = openpyxl.Workbook()
        sheet = workbook.active
        sheet.title = 'gallery_images'
        sheet.append(['filename', 'title', 'tags', 'description'])
        sheet.append(['../images/a.webp', '原有标题', '类型:访谈', '原有说明'])
        sheet.append(['../images/b.webp', '', None, ''])
        sheet.append(['../images/thumbs/c.webp', '不显示', '', ''])
        other = workbook.create_sheet('other')
        other.append(['原有内容', '=1+2'])
        workbook.save(self.path)
        workbook.close()
        self.store = GalleryStore(self.root)

    def test_save_only_tags_preserves_other_edits_and_creates_backup(self):
        self.assertEqual(len(self.store.entries), 2)
        workbook = openpyxl.load_workbook(self.path)
        workbook['other']['A1'] = '面板打开后的外部编辑'
        workbook.save(self.path)
        workbook.close()
        before = self.path.read_bytes()
        backup = self.store.save(self.store.entries[0], '类型:合照, 人物:高尾奏音')
        self.assertEqual(backup.read_bytes(), before)
        workbook = openpyxl.load_workbook(self.path)
        self.assertEqual(workbook['gallery_images']['C2'].value, '类型:合照, 人物:高尾奏音')
        self.assertEqual(workbook['gallery_images']['B2'].value, '原有标题')
        self.assertEqual(workbook['gallery_images']['D2'].value, '原有说明')
        self.assertEqual(workbook['other']['A1'].value, '面板打开后的外部编辑')
        self.assertEqual(workbook['other']['B1'].value, '=1+2')
        workbook.close()
        saved = self.path.read_bytes()
        self.assertIsNone(self.store.save(self.store.entries[0], self.store.entries[0]['tags']))
        self.assertEqual(self.path.read_bytes(), saved)

    def test_external_tag_edit_is_not_overwritten(self):
        workbook = openpyxl.load_workbook(self.path)
        workbook['gallery_images']['C2'] = '类型:日常'
        workbook.save(self.path)
        workbook.close()
        before = self.path.read_bytes()
        with self.assertRaisesRegex(ValueError, '其他操作修改'):
            self.store.save(self.store.entries[0], '类型:合照')
        self.assertEqual(self.path.read_bytes(), before)

    def test_failed_replace_keeps_original_and_draft_state(self):
        before = self.path.read_bytes()
        with patch('gallery_editor.os.replace', side_effect=PermissionError('Excel 占用')):
            with self.assertRaises(PermissionError):
                self.store.save(self.store.entries[0], '类型:合照')
        self.assertEqual(self.path.read_bytes(), before)
        self.assertEqual(self.store.entries[0]['tags'], '类型:访谈')
        self.assertFalse(list(self.path.parent.glob('.gallery-tags-*')))

    def test_excel_lock_and_duplicate_rows_are_rejected(self):
        lock = self.path.parent / '~$data.xlsx'
        lock.touch()
        with self.assertRaises(PermissionError):
            self.store.save(self.store.entries[0], '')
        lock.unlink()
        workbook = openpyxl.load_workbook(self.path)
        workbook['gallery_images'].append(['images/A.webp', '', '', ''])
        workbook.save(self.path)
        workbook.close()
        with self.assertRaisesRegex(ValueError, '重复'):
            self.store.save(self.store.entries[0], '')

    def test_clear_tags_and_invalid_input(self):
        with self.assertRaises(ValueError):
            self.store.save(self.store.entries[0], '类型：合照')
        self.store.save(self.store.entries[0], '')
        workbook = openpyxl.load_workbook(self.path)
        self.assertIsNone(workbook['gallery_images']['C2'].value)
        workbook.close()

    def test_batch_saves_multiple_images_with_one_workbook_load(self):
        real_load = openpyxl.load_workbook
        with patch('gallery_editor.openpyxl.load_workbook', wraps=real_load) as load:
            self.store.save_many([(self.store.entries[0], '类型:合照'),
                                  (self.store.entries[1], '人物:高尾奏音')])
            self.assertEqual(load.call_count, 1)
        saved = GalleryStore(self.root)
        self.assertEqual([entry['tags'] for entry in saved.entries], ['类型:合照', '人物:高尾奏音'])
        self.assertEqual(len(list((self.root / '_backups').glob('*.xlsx'))), 1)

    def test_batch_conflict_does_not_partially_write(self):
        before = self.path.read_bytes()
        conflicting = dict(self.store.entries[1], tags='外部修改')
        with self.assertRaises(ValueError):
            self.store.save_many([(self.store.entries[0], '类型:合照'), (conflicting, '类型:日常')])
        self.assertEqual(self.path.read_bytes(), before)


if __name__ == '__main__':
    unittest.main()
