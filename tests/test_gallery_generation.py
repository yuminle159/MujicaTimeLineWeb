import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

import openpyxl
import generate_all as generator


class GalleryGenerationTests(unittest.TestCase):
    def test_repeated_scan_preserves_manual_data_and_ignores_subfolders(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            images = root / 'images'
            images.mkdir()
            for name in ['a.webp', 'b.webp', 'c.webp', 'd.webp', 'new.webp']:
                (images / name).touch()
            (images / 'thumbs').mkdir()
            (images / 'thumbs/hidden.webp').touch()
            (images / 'fake.webp').mkdir()
            workbook = openpyxl.Workbook()
            sheet = workbook.active
            sheet.title = 'gallery_images'
            sheet.append(['filename', 'title', 'tags', 'description'])
            for name in ['../images/a.webp', 'images/b.webp', r'..\images\C.webp', 'd.webp']:
                sheet.append([name, '手填标题', '人物:高尾奏音, 类型:合照', '手填说明'])
            before = list(sheet.values)
            path = root / 'data.xlsx'
            workbook.save(path)
            workbook.close()
            with patch.object(generator, 'ROOT', str(root)), patch.object(generator, 'XLSX_PATH', str(path)):
                self.assertEqual(generator.scan_gallery_images(lambda _: None), 1)
                first_bytes = path.read_bytes()
                self.assertEqual(generator.scan_gallery_images(lambda _: None), 0)
                self.assertEqual(path.read_bytes(), first_bytes)
            workbook = openpyxl.load_workbook(path)
            values = list(workbook['gallery_images'].values)
            self.assertEqual(values[:-1], before)
            self.assertEqual(values[-1], ('../images/new.webp', None, None, None))
            workbook.close()

    def test_generation_and_search_use_same_unique_rows_without_mutating_tags(self):
        workbook = openpyxl.Workbook()
        sheet = workbook.active
        sheet.title = 'gallery_images'
        sheet.append(['filename', 'tags'])
        sheet.append(['../images/a.webp', '人物:高尾奏音'])
        sheet.append(['images/A.webp', '类型:合照'])
        sheet.append(['../images/thumbs/a.webp', '类型:缩略图'])
        before = list(sheet.values)
        with tempfile.TemporaryDirectory() as directory:
            output = Path(directory) / 'data.js'
            with patch.object(generator, 'ROOT', directory), patch.dict(generator.OUTPUTS, gallery=str(output)):
                for _ in range(2):
                    self.assertEqual(generator.generate_gallery(workbook), 1)
                    result = output.read_text(encoding='utf-8')
                    self.assertIn('高尾奏音', result)
                    self.assertNotIn('合照', result)
                    self.assertEqual(result.count('filename:'), 1)
        self.assertEqual(len(list(generator.gallery_rows(workbook))), 1)
        self.assertEqual(list(sheet.values), before)
        workbook.close()


if __name__ == '__main__':
    unittest.main()
