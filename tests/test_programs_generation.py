import json
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

import openpyxl
from openpyxl.comments import Comment
from openpyxl.styles import PatternFill

import generate_all
import sync_programs


class ProgramsTests(unittest.TestCase):
    def test_independent_cast_aliases_dates_and_clips(self):
        wb = openpyxl.Workbook()
        sheet = sync_programs.create_sheet(wb)
        sheet.append(['episode-a', 'Series', '#1', '2026/09/21', '20:00', '+09:00', '',
                      'Alias|Canonical|Unknown', 'https://example.com/a|https://example.com/b',
                      'https://example.com/full', 'images/programs/a.webp'])
        sheet['N7'], sheet['O7'], sheet['P7'], sheet['Q7'] = 'Canonical', 'Ave Mujica', '#881144', 'Alias'
        with tempfile.TemporaryDirectory() as directory:
            output = Path(directory) / 'data.js'
            with patch.dict(generate_all.OUTPUTS, programs=str(output)):
                self.assertEqual(generate_all.generate_programs(wb), 1)
            item = json.loads(output.read_text(encoding='utf-8').split(' = ', 1)[1].rstrip(';\n'))[0]
        self.assertEqual([p['name'] for p in item['performers']], ['Canonical', 'Unknown'])
        self.assertEqual(item['performers'][0]['bands'], ['Ave Mujica'])
        self.assertIsInstance(item['performers'][0]['order'], int)
        self.assertEqual(item['performers'][1]['bands'], [])
        self.assertEqual(item['date'], '2026-09-21')
        self.assertEqual(len(item['clips']), 2)
        self.assertEqual(item['cover'], '../images/programs/a.webp')
        self.assertEqual(sheet['I2'].value, 'https://example.com/a|https://example.com/b')
        wb.close()

    def test_save_preserves_empty_and_user_clip_cells(self):
        from hashlib import sha256
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / '_data' / 'data.xlsx'
            path.parent.mkdir()
            wb = openpyxl.Workbook()
            sheet = sync_programs.create_sheet(wb)
            sheet['I2'] = ''
            sheet['I3'] = 'https://example.com/clip'
            sheet['I3'].hyperlink = sheet['I3'].value
            sheet['I3'].comment = Comment('keep me', 'user')
            sheet['I3'].fill = PatternFill('solid', fgColor='FF881144')
            wb.save(path)
            sync_programs.save_workbook(wb, path, sha256(path.read_bytes()).hexdigest())
            restored = openpyxl.load_workbook(path)
            self.assertEqual(restored['programs']['I2'].value, '')
            self.assertEqual(restored['programs']['I2'].data_type, 's')
            self.assertEqual(restored['programs']['I3'].hyperlink.target, 'https://example.com/clip')
            self.assertEqual(restored['programs']['I3'].comment.text, 'keep me')
            self.assertEqual(restored['programs']['I3'].fill.fgColor.rgb, 'FF881144')
            restored.close()
            wb.close()


if __name__ == '__main__':
    unittest.main()
