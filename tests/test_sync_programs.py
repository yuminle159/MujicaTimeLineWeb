import unittest
from datetime import datetime
from pathlib import Path
import tempfile

import openpyxl
from openpyxl.styles import PatternFill

import sync_programs as sync


class ProgramImportTests(unittest.TestCase):
    def test_import_preserves_cast_colors_and_never_overwrites_curated_fields(self):
        with tempfile.TemporaryDirectory() as directory:
            source = Path(directory) / "source.xlsx"
            original = openpyxl.Workbook()
            sheet = original.active
            sheet["C1"] = "回数"
            sheet["N1"] = "出演者"
            sheet["B2"] = datetime(2024, 1, 4)
            sheet["C2"] = "#199"
            sheet["D2"] = "声优甲"
            sheet["E2"] = "嘉宾乙"
            sheet["N2"] = "声优甲"
            sheet["O2"] = "乐队甲"
            sheet["O2"].fill = PatternFill("solid", fgColor="FF881144")
            original.save(source)
            original.close()
            target = openpyxl.Workbook()
            target.active["A1"] = "existing data"
            sync.import_source(target, source)
            imported = target["programs"]
            self.assertEqual(imported["H2"].value, "声优甲|嘉宾乙")
            self.assertEqual(imported["P2"].value, "#881144")
            self.assertEqual(imported["E2"].value, "")
            self.assertIsNone(imported["H2"].fill.patternType)
            imported["I2"] = "https://example.com/clip1|https://example.com/clip2"
            imported["G2"] = "Manual title"
            report = sync.import_source(target, source)
            self.assertEqual(report["added_episodes"], 0)
            self.assertEqual(imported["G2"].value, "Manual title")
            self.assertIn("clip2", imported["I2"].value)
            self.assertEqual(target.active["A1"].value, "existing data")
            self.assertEqual(len(sync.read_episodes(imported)), 1)
            self.assertEqual(len(sync.read_cast(imported)[0]), 2)
            target.close()

    def test_duplicate_video_matches_remain_ambiguous_and_wrong_year_is_excluded(self):
        entry = {"program": "バンドリ！TV LIVE", "episode": "#199", "date": "2024/01/04", "title": "", "video_url": ""}
        candidates = [{"id": "aaaaaaaaaaa", "title": "バンドリ！TV LIVE 2024 #199"},
                      {"id": "bbbbbbbbbbb", "title": "バンドリ！TV LIVE 2023 #199"},
                      {"id": "ccccccccccc", "title": "バンドリ！TV LIVE 2024 #199 replay"}]
        matches, _ = sync.match_video(entry, candidates)
        self.assertEqual([v["id"] for v in matches], ["aaaaaaaaaaa", "ccccccccccc"])
        entry["video_url"] = "https://youtu.be/bbbbbbbbbbb"
        matches, mode = sync.match_video(entry, candidates)
        self.assertEqual(mode, "manual")
        self.assertEqual(matches[0]["id"], "bbbbbbbbbbb")

    def test_new_and_old_playlist_layouts_and_recommendation_scope(self):
        video = {"lockupViewModel": {"contentId": "aaaaaaaaaaa", "contentType": "LOCKUP_CONTENT_TYPE_VIDEO",
                 "metadata": {"lockupMetadataViewModel": {"title": {"content": "episode"}}}}}
        listing = {"itemSectionRenderer": {"contents": [video, {"continuationCommand": {"token": "playlist-next"}}]}}
        data = {"tabRenderer": {"selected": True, "content": {"sectionListRenderer": {"contents": [listing,
                {"continuationCommand": {"token": "recommendations-next"}}]}}}}
        videos, continuation = sync.parse_playlist_page(data)
        self.assertEqual(len(videos), 1)
        self.assertEqual(continuation, "playlist-next")
        response = {"contents": {"tabRenderer": {"selected": True}},
                    "onResponseReceivedActions": [{"appendContinuationItemsAction": {"continuationItems": [video]}}]}
        self.assertEqual(len(sync.parse_playlist_page(response)[0]), 1)
        old = {"playlistVideoRenderer": {"videoId": "bbbbbbbbbbb", "title": {"runs": [{"text": "old episode"}]}}}
        self.assertEqual(sync.parse_playlist_page(old)[0][0]["title"], "old episode")

    def test_special_title_may_omit_only_franchise_prefix(self):
        entry = {"program": "特别节目", "episode": "", "date": "2024/07/04", "title": "夏の大発表会2024", "video_url": ""}
        videos = [{"id": "aaaaaaaaaaa", "title": "バンドリ！　夏の大発表会2024"},
                  {"id": "bbbbbbbbbbb", "title": "バンドリ！　夏の大発表会2025"}]
        self.assertEqual(sync.match_video(entry, videos)[0], [videos[0]])


if __name__ == "__main__":
    unittest.main()
