"""Evidence-backed program updates; clips are a protected, read-only field."""
from __future__ import annotations

import argparse
from datetime import datetime
from hashlib import sha256
import importlib.util
import json
from pathlib import Path
import re
import sys

import openpyxl

EPISODE_FIELDS = {"program", "episode", "date", "time", "timezone", "title", "performers", "video_url", "cover", "notes"}
CAST_FIELDS = {"band", "band_color", "aliases"}


def load_project(root):
    script = root / "sync_programs.py"
    if not script.is_file() or not (root / "_data/data.xlsx").is_file():
        raise ValueError("项目缺少 sync_programs.py 或 _data/data.xlsx")
    spec = importlib.util.spec_from_file_location("programs_sync_skill", script)
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


def cell_state(cell):
    return {
        "value": cell.value, "type": cell.data_type,
        "hyperlink": {key: getattr(cell.hyperlink, key) for key in ("target", "location", "display", "tooltip")} if cell.hyperlink else None,
        "comment": {"text": cell.comment.text, "author": cell.comment.author,
                    "width": cell.comment.width, "height": cell.comment.height} if cell.comment else None,
        "style": list(cell._style) if cell.has_style else None,
    }


def clip_guard(sheet, module):
    col = module.columns(sheet)["clips"]
    return {e["program_id"]: cell_state(sheet.cell(e["row"], col)) for e in module.read_episodes(sheet)}


def check_clips(sheet, module, before):
    after = clip_guard(sheet, module)
    for ident, value in before.items():
        if ident not in after or after[ident] != value:
            raise ValueError("切片保护校验失败：" + ident)
    for ident in after.keys() - before.keys():
        if after[ident]["value"] is not None:
            raise ValueError("新节目不得自动填写切片：" + ident)


def validate_sheet(sheet, module):
    entries = module.read_episodes(sheet)
    cast, aliases = module.read_cast(sheet)
    numbers, urls = {}, {}
    for e in entries:
        key = (module.normalize(e["program"]), module.normalize(e["episode"]))
        if e["episode"]:
            if key in numbers:
                raise ValueError(f"节目期数重复：{numbers[key]} / {e['program_id']}")
            numbers[key] = e["program_id"]
        ident = module.video_id(e["video_url"])
        if ident:
            if ident in urls:
                raise ValueError(f"原视频重复：{urls[ident]} / {e['program_id']}")
            urls[ident] = e["program_id"]
        if e["time"]:
            if not e["date"]:
                raise ValueError("已填时间的节目必须有日期：" + e["program_id"])
            if not re.fullmatch(r"[+-](?:0\d|1[0-4]):[0-5]\d", e["timezone"]):
                raise ValueError("时区格式错误：" + e["program_id"])
    return entries, cast, aliases


def snapshot(root, module):
    path = root / "_data/data.xlsx"
    workbook = openpyxl.load_workbook(path)
    try:
        sheet = workbook["programs"]
        entries, cast, aliases = validate_sheet(sheet, module)
        return {
            "created_at": datetime.now().isoformat(timespec="seconds"),
            "workbook_sha256": sha256(path.read_bytes()).hexdigest(),
            "episodes": entries, "cast": cast, "clips_guard": clip_guard(sheet, module),
            "missing_fields": {e["program_id"]: [k for k in ("date", "time", "performers", "video_url", "cover") if not e[k]] for e in entries if any(not e[k] for k in ("date", "time", "performers", "video_url", "cover"))},
            "unmapped_performers": sorted({name for e in entries for name in module.split_values(e["performers"]) if not cast.get(aliases.get(name), {}).get("band")}),
        }
    finally:
        workbook.close()


def require_sources(sources, field=""):
    if not isinstance(sources, list) or not sources:
        raise ValueError("修改字段缺少来源：" + field)
    for url in sources:
        if not isinstance(url, str) or not (re.match(r"https?://[^\s/]+", url) or (field == "band_color" and re.fullmatch(r"workbook:programs!P[2-9]\d*|workbook:programs!P1\d+", url))):
            raise ValueError("来源必须是网页 URL；色值也可引用字典单元格")


def validate_fields(item, allowed):
    fields = item.get("fields", {})
    if not isinstance(fields, dict) or not fields:
        raise ValueError("补丁 fields 不能为空")
    unknown = set(fields) - allowed
    if unknown:
        raise ValueError("禁止修改字段：" + ", ".join(sorted(unknown)))
    for field, value in fields.items():
        if not isinstance(value, str):
            raise ValueError("字段必须为字符串：" + field)
        if value.startswith("="):
            raise ValueError("资料字段不能写入 Excel 公式")
        require_sources(item.get("sources", {}).get(field), field)
    return fields


def apply_patch(root, module, patch, dry_run=False):
    path = root / "_data/data.xlsx"
    digest = sha256(path.read_bytes()).hexdigest()
    if patch.get("expected_sha256") != digest:
        raise ValueError("工作簿已变更或缺少 expected_sha256，请重新快照并合并修改")
    if set(patch) - {"expected_sha256", "episodes", "cast", "cast_merges"}:
        raise ValueError("补丁含未知操作")
    workbook = openpyxl.load_workbook(path)
    changes = []
    try:
        sheet = workbook["programs"]
        cols = module.columns(sheet)
        before = clip_guard(sheet, module)

        def write(row, field, value, sources, ident):
            cell = sheet.cell(row, cols[field])
            if ("" if cell.value is None else str(cell.value)) != value:
                changes.append({"id": ident, "field": field, "before": cell.value, "after": value, "sources": sources})
                cell.value = value or None
                if field == "video_url" and cell.hyperlink:
                    cell.hyperlink = value or None

        def cast_rows():
            return {module.text(sheet.cell(r, cols["cast_name"]).value): r for r in range(2, sheet.max_row + 1) if sheet.cell(r, cols["cast_name"]).value}

        for item in patch.get("cast", []):
            fields = validate_fields(item, CAST_FIELDS)
            name = module.text(item.get("cast_name"))
            if not name or name.startswith("=") or "|" in name:
                raise ValueError("出演者标准名不合法")
            rows = cast_rows()
            row = rows.get(name, max(rows.values(), default=1) + 1)
            if name not in rows:
                write(row, "cast_name", name, list(item["sources"].values())[0], name)
            for field, value in fields.items():
                write(row, field, value, item["sources"][field], name)

        for merge in patch.get("cast_merges", []):
            require_sources(merge.get("sources"), "cast_merge")
            source, target = merge.get("from"), merge.get("into")
            rows = cast_rows()
            if source == target or source not in rows or target not in rows:
                raise ValueError("合并的两个姓名必须是不同的已有字典项")
            src, dest = rows[source], rows[target]
            for field in ("band", "band_color"):
                old, new = module.text(sheet.cell(src, cols[field]).value), module.text(sheet.cell(dest, cols[field]).value)
                if old and new and old != new:
                    raise ValueError("字典归属或色值冲突，请核实后修正再合并")
                if old and not new:
                    write(dest, field, old, merge["sources"], target)
            names = module.split_values(sheet.cell(dest, cols["aliases"]).value) + [source] + module.split_values(sheet.cell(src, cols["aliases"]).value)
            write(dest, "aliases", "|".join(dict.fromkeys(n for n in names if n != target)), merge["sources"], target)
            for field in ("cast_name", "band", "band_color", "aliases"):
                write(src, field, "", merge["sources"], source)

        _, aliases = module.read_cast(sheet)
        for item in patch.get("episodes", []):
            fields = validate_fields(item, EPISODE_FIELDS)
            if "performers" in fields:
                fields = dict(fields)
                fields["performers"] = "|".join(dict.fromkeys(aliases.get(n, n) for n in module.split_values(fields["performers"])))
            entries = {e["program_id"]: e for e in module.read_episodes(sheet)}
            ident = module.text(item.get("program_id"))
            if ident and ident in entries:
                row = entries[ident]["row"]
            else:
                if not fields.get("program") or not (fields.get("episode") or fields.get("title")):
                    raise ValueError("新节目需要节目系列，以及期数或标题")
                if not ident:
                    ident = module.stable_id(fields["program"], fields.get("episode", ""), module.date_text(fields.get("date", "")), fields.get("title", ""))
                if ident in entries:
                    raise ValueError("新节目已存在，请明确指定其 program_id 更新")
                row = max((e["row"] for e in entries.values()), default=1) + 1
                write(row, "program_id", ident, [], ident)
            for field, value in fields.items():
                if field == "date":
                    value = module.date_text(value)
                if field == "cover" and value:
                    cover = (root / value).resolve()
                    if not cover.is_relative_to((root / "images/programs").resolve()) or not cover.is_file():
                        raise ValueError("封面必须是 images/programs 下的已存在文件")
                if field == "video_url" and value and not re.match(r"https?://[^\s]+$", value):
                    raise ValueError("原视频链接格式错误")
                write(row, field, value, item["sources"][field], ident)

        entries, _, _ = validate_sheet(sheet, module)
        check_clips(sheet, module, before)
        sheet.auto_filter.ref = f"A1:L{max((e['row'] for e in entries), default=1)}"
        result = {"dry_run": dry_run, "changes": changes, "changed_episodes": len({c["id"] for c in changes if c["field"] in EPISODE_FIELDS | {"program_id"}}), "clips_preserved": True}
        if changes and not dry_run:
            result["backup"] = module.save_workbook(workbook, path, digest)
            saved = openpyxl.load_workbook(path)
            try:
                check_clips(saved["programs"], module, before)
            finally:
                saved.close()
        return result
    finally:
        workbook.close()


def collect(root, module):
    existing = snapshot(root, module)
    known = {module.video_id(e["video_url"]) for e in existing["episodes"]}
    with module.session() as client:
        videos = module.fetch_playlist(client, module.PLAYLIST, root / "_reports/programs-playlist.json")
    return {"fetched_at": datetime.now().isoformat(timespec="seconds"), "playlist": module.PLAYLIST,
            "videos": videos, "unregistered_candidates": [v for v in videos if v["id"] not in known],
            "scope_start": min((e["date"] for e in existing["episodes"] if e["date"]), default=None),
            "note": "候选项尚未判定日期和节目类型；不能全部直接导入。"}


def metadata(module, ids):
    result = {"videos": [], "failed": []}
    with module.session() as client:
        for ident in dict.fromkeys(ids):
            if not re.fullmatch(r"[A-Za-z0-9_-]{11}", ident):
                raise ValueError("无效视频 ID")
            url = "https://www.youtube.com/watch?v=" + ident
            try:
                response = client.get(url, timeout=30)
                response.raise_for_status()
                player = module.json_after(response.text, r"(?:var\s+)?ytInitialPlayerResponse\s*=\s*")
                details = player.get("videoDetails", {})
                if details.get("videoId") != ident or not details.get("title"):
                    raise ValueError("未取得指定视频元数据；不要推测缺失内容")
                micro = player.get("microformat", {}).get("playerMicroformatRenderer", {})
                result["videos"].append({
                    "id": ident, "url": url, "title": details["title"],
                    "description": details.get("shortDescription", ""),
                    "channel_id": details.get("channelId"), "author": details.get("author"),
                    "live_broadcast_details": micro.get("liveBroadcastDetails", {}),
                    "is_upcoming": details.get("isUpcoming"),
                    "publish_date_not_broadcast_date": micro.get("publishDate"),
                    "fetched_at": datetime.now().isoformat(timespec="seconds"),
                })
            except (ValueError, OSError, module.requests.RequestException) as error:
                result["failed"].append({"id": ident, "url": url, "error": str(error)})
    return result


def verify(root, module, original):
    current = snapshot(root, module)
    workbook = openpyxl.load_workbook(root / "_data/data.xlsx")
    try:
        check_clips(workbook["programs"], module, original["clips_guard"])
    finally:
        workbook.close()
    missing_covers = []
    for e in current["episodes"]:
        if e["cover"] and not (root / e["cover"]).is_file():
            missing_covers.append(e["program_id"])
    return {"clips_preserved": True, "episodes": len(current["episodes"]),
            "new_episodes": len(current["clips_guard"].keys() - original["clips_guard"].keys()),
            "missing_fields": current["missing_fields"], "unmapped_performers": current["unmapped_performers"],
            "missing_cover_files": missing_covers}


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("command", choices=("snapshot", "collect", "metadata", "apply", "verify"))
    parser.add_argument("--project", type=Path, default=Path.cwd())
    parser.add_argument("--output", type=Path, required=True)
    parser.add_argument("--patch", type=Path)
    parser.add_argument("--snapshot", type=Path)
    parser.add_argument("--video-id", action="append", default=[])
    parser.add_argument("--dry-run", action="store_true")
    args = parser.parse_args()
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8")
    root = args.project.resolve()
    module = load_project(root)
    if args.command == "snapshot":
        result = snapshot(root, module)
    elif args.command == "collect":
        result = collect(root, module)
    elif args.command == "metadata":
        if not args.video_id:
            parser.error("metadata 需要 --video-id")
        result = metadata(module, args.video_id)
    elif args.command == "apply":
        if not args.patch:
            parser.error("apply 需要 --patch")
        result = apply_patch(root, module, json.loads(args.patch.read_text(encoding="utf-8")), args.dry_run)
    else:
        if not args.snapshot:
            parser.error("verify 需要 --snapshot")
        result = verify(root, module, json.loads(args.snapshot.read_text(encoding="utf-8")))
    args.output.parent.mkdir(parents=True, exist_ok=True)
    args.output.write_text(json.dumps(result, ensure_ascii=False, indent=2, default=str), encoding="utf-8")
    print(f"{args.command} 完成：{args.output}")
    if result.get("failed") or result.get("missing_cover_files"):
        return 1
    return 0


if __name__ == "__main__":
    try:
        raise SystemExit(main())
    except (ValueError, RuntimeError, OSError, KeyError) as error:
        print(f"失败：{error}", file=sys.stderr)
        raise SystemExit(1)
