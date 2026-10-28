"""Append clips from the user's Bilibili season without rebuilding the workbook."""

import argparse
import copy
import json
import os
import re
import tempfile
import zipfile
from collections import defaultdict
from datetime import datetime
from hashlib import sha256
from pathlib import Path
from xml.sax.saxutils import escape
import xml.etree.ElementTree as ET

import openpyxl
import requests

MID = '3690988597938228'
SEASON = '7952732'
SOURCE = f'https://space.bilibili.com/{MID}/lists/{SEASON}?type=season'
API = 'https://api.bilibili.com/x/polymer/web-space/seasons_archives_list'
NS = {'m': 'http://schemas.openxmlformats.org/spreadsheetml/2006/main'}


def fetch():
    videos = []
    page = 1
    while True:
        response = requests.get(API, params={
            'mid': MID, 'season_id': SEASON, 'sort_reverse': 'false',
            'page_num': page, 'page_size': 100,
        }, headers={'User-Agent': 'Mozilla/5.0', 'Referer': f'https://space.bilibili.com/{MID}/'}, timeout=30)
        response.raise_for_status()
        payload = response.json()
        if payload.get('code') != 0:
            raise RuntimeError(f'Bilibili API: {payload.get("code")} {payload.get("message")}')
        data = payload['data']
        videos.extend(data['archives'])
        total = data['page']['total']
        if len(videos) >= total:
            break
        if not data['archives']:
            raise RuntimeError('Bilibili returned an incomplete season')
        page += 1
    if len(videos) != total or len({v['bvid'] for v in videos}) != total:
        raise RuntimeError('Season pages are incomplete or duplicated')
    return videos


def build_changes(sheet, videos):
    episodes = defaultdict(list)
    for row in range(2, sheet.max_row + 1):
        if sheet.cell(row, 2).value == 'バンドリ！TV LIVE':
            episodes[str(sheet.cell(row, 3).value)].append(row)
    by_row = defaultdict(list)
    unmatched = []
    already_present = 0
    for video in videos:
        title = video['title']
        if '邦TV' not in title:
            unmatched.append({'title': title, 'reason': 'missing 邦TV'})
            continue
        numbers = list(dict.fromkeys(re.findall(r'#(\d+)', title.split('邦TV', 1)[1])))
        if not numbers:
            unmatched.append({'title': title, 'reason': 'missing episode number'})
            continue
        for number in numbers:
            matches = episodes['#' + number]
            if len(matches) != 1:
                unmatched.append({'title': title, 'episode': '#' + number,
                                  'reason': f'{len(matches)} matching rows'})
                continue
            row = matches[0]
            url = 'https://www.bilibili.com/video/' + video['bvid']
            by_row[row].append((video['bvid'], title, url))
    changes = []
    labeled = 0
    for row, videos_for_row in by_row.items():
        old = sheet.cell(row, 9).value or ''
        entries = old.split('|') if old else []
        seen = set()
        urls = []
        for bvid, title, url in videos_for_row:
            if bvid in seen:
                continue
            seen.add(bvid)
            positions = [i for i, entry in enumerate(entries) if bvid in re.findall(r'BV[0-9A-Za-z]{10}', entry)]
            if positions:
                already_present += 1
                for i in positions:
                    if entries[i].strip() == url:
                        entries[i] = f'{title}=>{url}'
                        labeled += 1
            else:
                entries.append(f'{title}=>{url}')
                urls.append(url)
        new = '|'.join(entries)
        if new != old:
            changes.append({'row': row, 'program_id': sheet.cell(row, 1).value,
                            'episode': sheet.cell(row, 3).value, 'before': old,
                            'after': new, 'urls': urls})
    return changes, unmatched, already_present, labeled


def write_only_clip_cells(path, changes, digest):
    with zipfile.ZipFile(path) as original:
        book = ET.fromstring(original.read('xl/workbook.xml'))
        entry = next(s for s in book.find('m:sheets', NS) if s.attrib['name'] == 'programs')
        rid = entry.attrib['{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id']
        relationship = next(s for s in ET.fromstring(original.read('xl/_rels/workbook.xml.rels')) if s.attrib['Id'] == rid)
        name = relationship.attrib['Target']
        name = name.lstrip('/') if name.startswith('/') else 'xl/' + name
        xml = original.read(name).decode('utf-8')
        for change in changes:
            address = 'I' + str(change['row'])
            pattern = r'<c\b(?=[^>]*\br="' + address + r'")[^>]*(?:/>|>.*?</c>)'
            match = re.search(pattern, xml, re.S)
            if not match:
                raise RuntimeError(f'Cannot find clip cell {address}')
            tag = re.sub(r'\s+t="[^"]*"', '', match.group().split('>', 1)[0].rstrip('/'))
            replacement = tag + ' t="inlineStr"><is><t xml:space="preserve">' + escape(change['after']) + '</t></is></c>'
            xml = xml[:match.start()] + replacement + xml[match.end():]
        ET.fromstring(xml)
        handle, temporary = tempfile.mkstemp(suffix='.xlsx', dir=path.parent)
        os.close(handle)
        try:
            with zipfile.ZipFile(temporary, 'w') as output:
                for item in original.infolist():
                    output.writestr(copy.copy(item), xml.encode('utf-8') if item.filename == name else original.read(item.filename))
            with zipfile.ZipFile(temporary) as staged:
                if staged.namelist() != original.namelist():
                    raise RuntimeError('Workbook package structure changed')
                if any(staged.read(n) != original.read(n) for n in original.namelist() if n != name):
                    raise RuntimeError('An unrelated workbook part changed')
            staged_book = openpyxl.load_workbook(temporary, read_only=True)
            try:
                for change in changes:
                    if staged_book['programs'].cell(change['row'], 9).value != change['after']:
                        raise RuntimeError(f'Clip cell I{change["row"]} did not save correctly')
            finally:
                staged_book.close()
            if sha256(path.read_bytes()).hexdigest() != digest:
                raise RuntimeError('Workbook changed during update; nothing written')
            backup = path.parent.parent / '_backups' / ('programs-clips-' + datetime.now().strftime('%Y%m%d-%H%M%S-%f')) / path.name
            backup.parent.mkdir(parents=True)
            backup.write_bytes(path.read_bytes())
            if sha256(path.read_bytes()).hexdigest() != digest:
                raise RuntimeError('Workbook changed after backup; nothing written')
            try:
                os.replace(temporary, path)
            except PermissionError:
                # Some Windows spreadsheet integrations permit writes but deny rename.
                # The original bytes are already backed up, and the staged ZIP was checked.
                if sha256(path.read_bytes()).hexdigest() != digest:
                    raise RuntimeError('Workbook changed before fallback write; nothing written')
                with open(path, 'r+b') as target, open(temporary, 'rb') as source:
                    target.seek(0)
                    target.write(source.read())
                    target.truncate()
                    target.flush()
                    os.fsync(target.fileno())
                if sha256(path.read_bytes()).hexdigest() != sha256(Path(temporary).read_bytes()).hexdigest():
                    raise RuntimeError(f'Fallback write failed; restore workbook from {backup}')
            return str(backup)
        finally:
            Path(temporary).unlink(missing_ok=True)


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--project', type=Path, required=True)
    parser.add_argument('--apply', action='store_true')
    parser.add_argument('--report', type=Path)
    args = parser.parse_args()
    path = args.project / '_data/data.xlsx'
    digest = sha256(path.read_bytes()).hexdigest()
    videos = fetch()
    book = openpyxl.load_workbook(path, read_only=True)
    try:
        changes, unmatched, existing, labeled = build_changes(book['programs'], videos)
    finally:
        book.close()
    report = {'source': SOURCE, 'workbook_sha256_before': digest,
              'videos': len(videos), 'added_links': sum(len(c['urls']) for c in changes),
              'updated_episodes': len(changes), 'existing_links': existing,
              'labeled_existing_links': labeled,
              'unmatched': unmatched, 'changes': changes}
    if args.apply and changes:
        report['backup'] = write_only_clip_cells(path, changes, digest)
        report['workbook_sha256_after'] = sha256(path.read_bytes()).hexdigest()
    if args.report:
        args.report.parent.mkdir(parents=True, exist_ok=True)
        args.report.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding='utf-8')
    print(json.dumps({k: v for k, v in report.items() if k != 'changes'}, ensure_ascii=False, indent=2))


if __name__ == '__main__':
    main()
